#!/usr/bin/env python3
"""Teste automático de ponta a ponta com os dados de exemplo (assets/sample). Não toca em dados reais.
Uso: python jh.py selftest
"""
import json, os, shutil, subprocess, sys, tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SAMPLE = HERE.parent / "assets" / "sample"
JH = [sys.executable, str(HERE / "jh.py")]
AGENTS = ["fact_checker", "ats", "hr", "hiring_manager", "language_editor", "skeptic"]
fails = []


def run(*args):
    r = subprocess.run(JH + list(args), capture_output=True, text=True, encoding="utf-8")
    return r.returncode, (r.stdout + r.stderr)


def expect(name, cond, extra=""):
    print(("PASS  " if cond else "FAIL  ") + name + (f"  [{extra}]" if (extra and not cond) else ""))
    if not cond:
        fails.append(name)


def verdict(job, rnd, ag, h, verdict="GREEN", score=None, evid=3, mtime=None):
    score = score if score is not None else (100 if ag == "fact_checker" else 95)
    p = job / "revisoes" / f"round-{rnd}" / f"{ag}.json"
    p.write_text(json.dumps({
        "agent": ag, "round": rnd, "artifact_hash": h, "verdict": verdict, "score": score,
        "blockers": [] if verdict == "GREEN" else [{"where": "cv.summary", "problem": "x", "fix_hint": "y"}],
        "warnings": [], "evidence": ["e"] * evid, "weakest_points": ["a", "b", "c"]}), encoding="utf-8")
    if mtime:
        os.utime(p, (mtime, mtime))


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    try:
        import docx  # noqa
    except ImportError:
        sys.exit("python-docx em falta: pip install python-docx")
    tmp = Path(tempfile.mkdtemp(prefix="jh_selftest_"))
    root = tmp / "candidaturas"
    try:
        (root / "perfil").mkdir(parents=True)
        master = json.loads((SAMPLE / "master_cv.json").read_text(encoding="utf-8"))
        master.pop("audit", None)
        for e in master["experience"]:
            e.pop("confirmed_at", None)
        mp = root / "perfil" / "master_cv.json"
        mp.write_text(json.dumps(master, ensure_ascii=False), encoding="utf-8")
        code, out = run("audit", "status", str(mp))
        expect("audit status: CV mestre sem auditoria bloqueia", code == 1 and "EM FALTA" in out, out)
        code, out = run("audit", "confirm", str(mp))
        expect("audit confirm recusa sem --confirmado-pela-pessoa", code != 0 and "audit" not in json.loads(mp.read_text(encoding="utf-8")), out)
        code, out = run("init", "--root", str(root), "--empresa", "Banco Exemplo, S.A.", "--cargo", "Analista Financeiro Sénior")
        job = Path(out.strip().splitlines()[-1])
        expect("init cria pasta da vaga com rascunho/ e revisoes/", (job / "rascunho").is_dir() and (job / "revisoes").is_dir())
        for f in ("cv.json", "carta.json"):
            shutil.copy(SAMPLE / f, job / "rascunho" / f)
        shutil.copy(SAMPLE / "keywords.txt", job / "keywords.txt")
        code, out = run("build", "cv", str(job / "rascunho/cv.json"), str(job / "rascunho/CV_Ana_Banco.docx"))
        expect("build cv", code == 0, out)
        code, out = run("build", "letter", str(job / "rascunho/carta.json"), str(job / "rascunho/Carta_Banco.docx"))
        expect("build carta", code == 0, out)
        (job / "revisoes" / "round-1").mkdir()
        code, h = run("hash", str(job), "--write"); h = h.strip()
        code, out = run("check", "--job", str(job), "--docx", str(job / "rascunho/CV_Ana_Banco.docx"), "--empresa", "Banco Exemplo")
        expect("ats_check passa com o CV de exemplo", code == 0, out)

        code, out = run("gate", str(job))
        expect("gate: sem auditoria do CV mestre recusa (4), mesmo com tudo o resto", code == 4, out)
        code, out = run("audit", "confirm", str(mp), "--confirmado-pela-pessoa", "--notes", "selftest")
        expect("audit confirm com confirmação explícita", code == 0, out)
        code, out = run("audit", "status", str(mp))
        expect("audit status: OK depois de confirmado", code == 0, out)
        code, out = run("gate", str(job))
        expect("gate: ronda sem veredictos = incompleta (3)", code == 3, out)
        for ag in AGENTS:
            verdict(job, 1, ag, h)
        code, out = run("gate", str(job))
        expect("gate: tudo GREEN = APROVADO (0) e cria final/", code == 0 and (job / "final").is_dir()
               and (job / "revisoes/APROVADO.json").exists(), out)

        # reprovações que têm de acontecer
        verdict(job, 1, "hr", h, evid=1)
        code, out = run("gate", str(job))
        expect("gate: GREEN sem provas é rejeitado (1) e final/ removida", code == 1 and not (job / "final").exists(), out)
        verdict(job, 1, "hr", h)
        verdict(job, 1, "fact_checker", h, score=96)
        code, out = run("gate", str(job))
        expect("gate: fact_checker abaixo de 100 é rejeitado", code == 1, out)
        verdict(job, 1, "fact_checker", h)
        verdict(job, 1, "ats", "0000000000000000")
        code, out = run("gate", str(job))
        expect("gate: veredicto de versão antiga é rejeitado", code == 1 and "DESATUALIZADA" in out + (job / "revisoes/round-1/feedback_para_escritor.md").read_text(encoding="utf-8"), out)
        verdict(job, 1, "ats", h)
        verdict(job, 1, "skeptic", h, mtime=1_000_000_000)  # gravado "antes" dos outros
        code, out = run("gate", str(job))
        expect("gate: skeptic gravado antes dos outros é rejeitado", code == 1, out)
        verdict(job, 1, "skeptic", h)
        code, out = run("gate", str(job))
        expect("gate: re-aprova com tudo válido", code == 0, out)

        # auditoria antiga expira
        m2 = json.loads(mp.read_text(encoding="utf-8")); m2["audit"]["date"] = "2020-01-01"
        mp.write_text(json.dumps(m2, ensure_ascii=False), encoding="utf-8")
        code, out = run("gate", str(job))
        expect("gate: auditoria com mais de 90 dias recusa (4) e remove final/", code == 4 and not (job / "final").exists(), out)
        m2["audit"]["date"] = json.loads(mp.read_text(encoding="utf-8"))["experience"][0]["confirmed_at"]
        m2["experience"][0]["role_note"] = "funções de direção financeira"
        mp.write_text(json.dumps(m2, ensure_ascii=False), encoding="utf-8")
        code, out = run("audit", "status", str(mp))
        expect("role_note sem confirmação bloqueia a auditoria", code == 1 and "role_note" in out, out)
        m2["experience"][0].pop("role_note"); mp.write_text(json.dumps(m2, ensure_ascii=False), encoding="utf-8")
        run("audit", "confirm", str(mp), "--confirmado-pela-pessoa")
        verdict(job, 1, "skeptic", h)  # restaurar ordem dos veredictos após as alterações acima

        # role_note aparece no .docx
        cvn = json.loads((job / "rascunho/cv.json").read_text(encoding="utf-8")); cvn["experience"][0]["role_note"] = "funções de reporting"
        (job / "rascunho/cv_note.json").write_text(json.dumps(cvn, ensure_ascii=False), encoding="utf-8")
        run("build", "cv", str(job / "rascunho/cv_note.json"), str(tmp / "note.docx"))
        from docx import Document
        expect("role_note aparece junto ao título no .docx", "(funções de reporting)" in "\n".join(p.text for p in Document(tmp / "note.docx").paragraphs))

        # alterar o CV invalida a aprovação
        cvp = job / "rascunho/cv.json"
        d = json.loads(cvp.read_text(encoding="utf-8")); d["summary"] += " Extra."
        cvp.write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")
        code, out = run("gate", str(job))
        expect("alterar o CV invalida a aprovação (final/ removida)", code != 0 and not (job / "final").exists(), out)

        # ats_check apanha problemas
        d["experience"][0]["bullets"] = ["Fez tarefas várias sem números."] * 4; d["summary"] = "Proativo e dinâmico. [CONFIRMAR]"
        cvp.write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")
        code, out = run("check", "--job", str(job))
        expect("ats_check apanha placeholder, falta de números e keywords", code == 1 and "sem_placeholders" in out, out)

        code, out = run("tracker", str(root), "list")
        expect("tracker list", code == 0 and "Banco Exemplo" in out, out)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print("\n" + ("TUDO OK" if not fails else f"{len(fails)} FALHA(S): " + "; ".join(fails)))
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
