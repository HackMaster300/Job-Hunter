#!/usr/bin/env python3
"""Extrai texto de um CV existente (.docx, .pdf, .txt, .md) para o agente o ler e montar o master_cv.json.

Uso: python extract_text.py ficheiro.docx > cv.txt
.pdf: usa pdftotext (poppler) se existir, senão pypdf se instalado. Se nenhum existir, ler o PDF diretamente
com a ferramenta de leitura do agente (o Claude Code lê PDFs nativamente).
"""
import shutil, subprocess, sys
from pathlib import Path


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    p = Path(sys.argv[1])
    if not p.exists():
        sys.exit(f"Não existe: {p}")
    ext = p.suffix.lower()
    if ext in (".txt", ".md"):
        print(p.read_text(encoding="utf-8", errors="replace"))
    elif ext == ".docx":
        from docx import Document
        d = Document(p)
        for para in d.paragraphs:
            if para.text.strip():
                print(para.text)
        for t in d.tables:  # CVs de terceiros usam muitas tabelas
            for row in t.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    print(" | ".join(cells))
    elif ext == ".pdf":
        if shutil.which("pdftotext"):
            print(subprocess.run(["pdftotext", "-layout", str(p), "-"], capture_output=True, text=True).stdout)
        else:
            try:
                from pypdf import PdfReader
                for pg in PdfReader(str(p)).pages:
                    print(pg.extract_text() or "")
            except ImportError:
                sys.exit("Sem pdftotext nem pypdf. Instalar (pip install pypdf) ou ler o PDF com a ferramenta de leitura do agente.")
    else:
        sys.exit(f"Formato não suportado: {ext}")


if __name__ == "__main__":
    main()
