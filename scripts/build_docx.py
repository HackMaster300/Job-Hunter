#!/usr/bin/env python3
"""Gera CV (ATS-friendly, coluna única) ou carta de apresentação em .docx a partir de JSON.

Uso:
  python build_docx.py cv     cv.json     saida.docx [--pdf]
  python build_docx.py letter carta.json  saida.docx [--pdf]
Idioma dos títulos: campo "lang" no JSON ("pt", "en", "fr", "es"; por omissão "pt"); qualquer título pode ser
substituído com "labels": {"experience": "..."}. Schemas em references/cv-schema.md.
Requer python-docx (pip install python-docx). --pdf usa LibreOffice se existir.
"""
import json, shutil, subprocess, sys
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

LABELS = {
    "pt": {"summary": "Resumo profissional", "skills": "Competências", "experience": "Experiência profissional",
           "projects": "Projetos", "education": "Formação", "certifications": "Certificações",
           "extras": "Informação adicional", "subject": "Assunto"},
    "en": {"summary": "Professional summary", "skills": "Skills", "experience": "Professional experience",
           "projects": "Projects", "education": "Education", "certifications": "Certifications",
           "extras": "Additional information", "subject": "Subject"},
    "fr": {"summary": "Résumé professionnel", "skills": "Compétences", "experience": "Expérience professionnelle",
           "projects": "Projets", "education": "Formation", "certifications": "Certifications",
           "extras": "Informations complémentaires", "subject": "Objet"},
    "es": {"summary": "Resumen profesional", "skills": "Competencias", "experience": "Experiencia profesional",
           "projects": "Proyectos", "education": "Formación", "certifications": "Certificaciones",
           "extras": "Información adicional", "subject": "Asunto"},
}


def labels(data):
    L = dict(LABELS.get(str(data.get("lang", "pt")).lower()[:2], LABELS["pt"]))
    L.update(data.get("labels") or {})
    return L


def validate_cv(d):
    errs = []
    if not isinstance(d, dict) or not d.get("name"):
        errs.append("falta 'name'")
        return errs
    for k in ("contacts", "certifications", "extras"):
        if k in d and not isinstance(d[k], list):
            errs.append(f"'{k}' deve ser lista")
    if "skills" in d and not isinstance(d["skills"], dict):
        errs.append("'skills' deve ser objeto {grupo: [itens]}")
    for i, e in enumerate(d.get("experience", [])):
        if not isinstance(e, dict) or not e.get("role"):
            errs.append(f"experience[{i}] sem 'role'")
        elif not isinstance(e.get("bullets", []), list):
            errs.append(f"experience[{i}].bullets deve ser lista")
    return errs


FONT = "Calibri"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)


def base_doc(margin_cm=1.8, size=10.5):
    d = Document()
    for s in d.sections:
        s.top_margin = s.bottom_margin = Cm(margin_cm)
        s.left_margin = s.right_margin = Cm(margin_cm)
    st = d.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(size)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.space_before = Pt(0)
    return d


def para(d, text="", bold=False, italic=False, size=None, align=None, after=0, before=0, color=None):
    p = d.add_paragraph()
    if text:
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
    if align == "center": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    return p


def bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    e = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "1F3A5F")):
        e.set(qn(f"w:{k}"), v)
    b.append(e)
    pPr.append(b)


def section(d, title):
    p = para(d, title.upper(), bold=True, size=11, before=9, after=3, color=ACCENT)
    bottom_border(p)


def bullet(d, text):
    p = d.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Cm(0.6)
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    return p


def line_lr(d, left, right, bold_left=True, italic_right=False, note=None):
    """Linha com texto à esquerda e à direita (tab alinhado à margem direita), sem tabelas."""
    p = d.add_paragraph()
    sec = d.sections[0]
    width = sec.page_width - sec.left_margin - sec.right_margin
    p.paragraph_format.tab_stops.add_tab_stop(width, alignment=2)  # 2 = RIGHT
    r = p.add_run(left); r.bold = bold_left
    if note:
        rn = p.add_run(f" ({note})"); rn.italic = True
    if right:
        r2 = p.add_run("\t" + right); r2.italic = italic_right
    p.paragraph_format.space_before = Pt(4)
    return p


def build_cv(data, out):
    L = labels(data)
    d = base_doc()
    para(d, data["name"], bold=True, size=20, align="center", color=ACCENT)
    if data.get("headline"):
        para(d, data["headline"], size=11.5, align="center", after=2)
    if data.get("contacts"):
        para(d, "  |  ".join(data["contacts"]), size=9.5, align="center", after=2)

    if data.get("summary"):
        section(d, L["summary"])
        para(d, data["summary"])

    if data.get("skills"):
        section(d, L["skills"])
        for group, items in data["skills"].items():
            p = d.add_paragraph()
            r = p.add_run(f"{group}: "); r.bold = True
            p.add_run(", ".join(items))
            p.paragraph_format.space_after = Pt(1)

    if data.get("experience"):
        section(d, L["experience"])
        for e in data["experience"]:
            line_lr(d, e.get("role", ""), e.get("period", ""), note=e.get("role_note"))
            sub = ", ".join(x for x in (e.get("company"), e.get("location")) if x)
            if sub:
                para(d, sub, italic=True, after=1)
            for b in e.get("bullets", []):
                bullet(d, b)

    if data.get("projects"):
        section(d, L["projects"])
        for pr in data["projects"]:
            p = d.add_paragraph()
            r = p.add_run(pr.get("name", "")); r.bold = True
            if pr.get("description"):
                p.add_run(" - " + pr["description"])
            p.paragraph_format.space_after = Pt(1)

    if data.get("education"):
        section(d, L["education"])
        for ed in data["education"]:
            line_lr(d, ed.get("degree", ""), ed.get("period", ""))
            if ed.get("school"):
                para(d, ed["school"], italic=True)
            if ed.get("details"):
                para(d, ed["details"], size=10)

    if data.get("certifications"):
        section(d, L["certifications"])
        for c in data["certifications"]:
            bullet(d, c)

    if data.get("extras"):
        section(d, L["extras"])
        for x in data["extras"]:
            bullet(d, x)

    d.save(out)


def build_letter(data, out):
    L = labels(data)
    d = base_doc(margin_cm=2.5, size=11)
    for i, l in enumerate(data.get("sender", [])):
        para(d, l, bold=(i == 0), size=12 if i == 0 else 10.5)
    if data.get("date"):
        para(d, data["date"], before=14, after=10)
    for l in data.get("recipient", []):
        para(d, l)
    if data.get("subject"):
        para(d, L["subject"] + ": " + data["subject"], bold=True, before=12, after=10)
    if data.get("salutation"):
        para(d, data["salutation"], after=8)
    for t in data.get("paragraphs", []):
        p = para(d, t, after=8)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if data.get("closing"):
        para(d, data["closing"], before=6, after=26)
    if data.get("signature"):
        para(d, data["signature"], bold=True)
    d.save(out)


def to_pdf(docx_path):
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        print("AVISO: LibreOffice não encontrado; guardar como PDF a partir do Word.", file=sys.stderr)
        return None
    out = Path(docx_path).parent
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(out), str(docx_path)],
                   capture_output=True, timeout=180)
    pdf = Path(docx_path).with_suffix(".pdf")
    return pdf if pdf.exists() else None


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    args = [a for a in sys.argv[1:] if a != "--pdf"]
    want_pdf = "--pdf" in sys.argv[1:]
    if len(args) != 3 or args[0] not in ("cv", "letter"):
        sys.exit(__doc__)
    kind, src, out = args
    with open(src, encoding="utf-8") as f:
        data = json.load(f)
    if kind == "cv":
        errs = validate_cv(data)
        if errs:
            sys.exit("cv.json inválido: " + "; ".join(errs))
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    (build_cv if kind == "cv" else build_letter)(data, out)
    print(out)
    if want_pdf:
        pdf = to_pdf(out)
        if pdf:
            print(pdf)


if __name__ == "__main__":
    main()
