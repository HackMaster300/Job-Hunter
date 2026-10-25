#!/usr/bin/env python3
"""Verificações mecânicas (determinísticas) do CV e da carta. Alimenta o agente `ats`.

Uso:
  python ats_check.py --job <pasta_vaga> [--docx rascunho/CV.docx] [--empresa "Empresa"] [--round N]
  (lê rascunho/cv.json, rascunho/carta.json, keywords.txt; grava revisoes/round-N/ats_check.json)

keywords.txt: uma por linha. "*" no início = OBRIGATÓRIA. Sinónimos com "|":
  *reconciliações bancárias|bank reconciliation
Sai com código 1 se alguma verificação bloqueante falhar.
"""
import argparse, json, math, re, subprocess, sys, tempfile, unicodedata
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from _common import load_config, utf8

PLACEHOLDERS = re.compile(r"\[(CONFIRMAR|TODO|XXX|[A-Z ]{3,})\]|\bTODO\b|\bXXX\b|lorem ipsum|\{\{.*?\}\}", re.I)
CLICHES = ["proativo", "dinamico", "orientado a resultados", "team player", "go-getter", "apaixonado por desafios",
           "venho por este meio", "hardworking", "results-driven", "self-starter", "detail-oriented", "proactive",
           "passionate about", "think outside the box"]
PRESENT = {"presente", "present", "atual", "atualidade", "current", "actualidade", "actual", "now"}
MONTHS = {m: i + 1 for i, ms in enumerate([
    ["jan", "janeiro", "january"], ["fev", "fevereiro", "feb", "february"], ["mar", "marco", "march"],
    ["abr", "abril", "apr", "april"], ["mai", "maio", "may"], ["jun", "junho", "june"], ["jul", "julho", "july"],
    ["ago", "agosto", "aug", "august"], ["set", "setembro", "sep", "sept", "september"],
    ["out", "outubro", "oct", "october"], ["nov", "novembro", "november"], ["dez", "dezembro", "dec", "december"]])
    for m in ms}


def norm(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", s)


def stem_word(w):
    """Stemming leve PT/EN para casar plurais: reconciliacoes/reconciliacao, reports/report."""
    for suf, rep in (("oes", "ao"), ("aes", "ao"), ("ais", "al"), ("eis", "el"), ("ies", "y")):
        if w.endswith(suf) and len(w) > 5:
            return w[: -len(suf)] + rep
    if w.endswith("s") and len(w) > 3 and not w.endswith("ss"):
        return w[:-1]
    return w


def stem_text(s):
    return " " + " ".join(stem_word(w) for w in re.findall(r"[a-z0-9+#.]+", norm(s))) + " "


def kw_present(stemmed_text, kw):
    for alt in kw.split("|"):
        k = stem_text(alt).strip()
        if k and f" {k} " in stemmed_text:
            return True
    return False


def cv_text(cv):
    parts = [cv.get("name", ""), cv.get("headline", ""), " ".join(cv.get("contacts", [])), cv.get("summary", "")]
    for g, items in (cv.get("skills") or {}).items():
        parts += [g] + list(items)
    for e in cv.get("experience", []):
        parts += [e.get("role", ""), e.get("role_note", ""), e.get("company", ""), e.get("location", ""), e.get("period", "")] + e.get("bullets", [])
    for e in cv.get("education", []):
        parts += [e.get("degree", ""), e.get("school", ""), e.get("period", ""), e.get("details", "")]
    parts += cv.get("certifications", []) + cv.get("extras", [])
    for p in cv.get("projects", []):
        parts += [p.get("name", ""), p.get("description", "")]
    return "\n".join(x for x in parts if x)


def top_text(cv):
    parts = [cv.get("headline", ""), cv.get("summary", "")]
    for g, items in (cv.get("skills") or {}).items():
        parts += list(items)
    if cv.get("experience"):
        e = cv["experience"][0]
        parts += [e.get("role", "")] + e.get("bullets", [])
    return "\n".join(parts)


def parse_period(s):
    """(início, fim) como (ano, mês) ou None."""
    toks = re.split(r"\s*[-–—]\s*|\s+a\s+|\s+to\s+", s.strip(), maxsplit=1)
    today = date.today()

    def one(t, end):
        t = norm(t).strip()
        if t in PRESENT:
            return (today.year, today.month)
        y = re.search(r"(19|20)\d{2}", t)
        if not y:
            return None
        m = next((MONTHS[w] for w in re.findall(r"[a-z]+", t) if w in MONTHS), None)
        return (int(y.group()), m or (12 if end else 1))

    if len(toks) == 1:
        a = one(toks[0], False)
        return (a, a) if a else None
    a, b = one(toks[0], False), one(toks[1], True)
    return (a, b) if a and b else None


def add(res, name, ok, detail, blocking=True):
    res["checks"].append({"name": name, "pass": bool(ok), "blocking": blocking, "detail": detail})


def check_cv(cv, kws, res, cfg):
    text = cv_text(cv)
    st, stop = stem_text(text), stem_text(top_text(cv))
    words = len(text.split())
    lo, hi = cfg["cv_words"]
    add(res, "tem_nome", bool(cv.get("name")), cv.get("name", ""))
    contacts = " ".join(cv.get("contacts", []))
    add(res, "email_presente", bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", contacts)), "email nos contactos")
    add(res, "telefone_presente", bool(re.search(r"\+?\d[\d\s().-]{7,}", contacts)), "telefone nos contactos")
    add(res, "tem_experiencia", bool(cv.get("experience")), f"{len(cv.get('experience', []))} cargos")
    add(res, "tem_formacao", bool(cv.get("education")), f"{len(cv.get('education', []))} entradas", blocking=False)
    m = PLACEHOLDERS.search(text)
    add(res, "sem_placeholders", not m, m.group() if m else "limpo")
    add(res, f"tamanho_palavras_{lo}_{hi}", lo <= words <= hi, f"{words} palavras")
    s = cv.get("summary", "")
    add(res, "resumo_ate_600_chars", 0 < len(s) <= 600, f"{len(s)} caracteres")

    bullets = [b for e in cv.get("experience", []) for b in e.get("bullets", [])]
    long_b = [b[:60] for b in bullets if len(b) > 240]
    add(res, "bullets_nao_longos", not long_b, f"{len(long_b)} acima de 240 caracteres: {long_b[:2]}")
    many = [e.get("role") for e in cv.get("experience", []) if len(e.get("bullets", [])) > 7]
    add(res, "max_7_bullets_por_cargo", not many, str(many) if many else "ok")
    quant = [b for b in bullets if re.search(r"\d", b)]
    ratio = len(quant) / len(bullets) if bullets else 0
    add(res, f"bullets_quantificados_>={cfg['quant_ratio']:.0%}", ratio >= cfg["quant_ratio"], f"{len(quant)}/{len(bullets)} ({ratio:.0%})")
    starts = [norm(b).split(" ")[0] for b in bullets if b]
    rep = sorted({w for w in starts if starts.count(w) > 2})
    add(res, "verbos_iniciais_nao_repetidos", not rep, f"repetidos >2x: {rep}", blocking=False)
    cl = [c for c in CLICHES if c in norm(text)]
    add(res, "sem_cliches", not cl, f"clichés: {cl}", blocking=False)

    periods = [(e.get("role"), parse_period(e.get("period", ""))) for e in cv.get("experience", [])]
    bad = [r for r, p in periods if p is None]
    add(res, "datas_legiveis", not bad, f"períodos não reconhecidos: {bad}" if bad else "ok")
    ok_p = [(r, p) for r, p in periods if p]
    order_ok = all(ok_p[i][1][0] >= ok_p[i + 1][1][0] for i in range(len(ok_p) - 1))
    add(res, "ordem_cronologica_inversa", order_ok, "mais recente primeiro")
    gaps = []
    for i in range(len(ok_p) - 1):
        a, b = ok_p[i][1][0], ok_p[i + 1][1][1]
        months = (a[0] - b[0]) * 12 + (a[1] - b[1])
        if months > 6:
            gaps.append(f"{ok_p[i + 1][0]} -> {ok_p[i][0]}: ~{months} meses")
    add(res, "sem_lacunas_>6_meses", not gaps, "; ".join(gaps) if gaps else "ok", blocking=False)

    req = [k for k, r in kws if r]
    opt = [k for k, r in kws if not r]
    miss_req = [k for k in req if not kw_present(st, k)]
    miss_opt = [k for k in opt if not kw_present(st, k)]
    cov = (len(req) - len(miss_req)) / len(req) if req else 1
    res["keywords"] = {"obrigatorias_total": len(req), "obrigatorias_em_falta": miss_req,
                       "desejaveis_em_falta": miss_opt, "cobertura_obrigatorias": round(cov, 2),
                       "keywords_no_topo": [k for k, _ in kws if kw_present(stop, k)]}
    add(res, f"keywords_obrigatorias>={cfg['kw_coverage']:.0%}", cov >= cfg["kw_coverage"],
        f"{cov:.0%}; em falta: {miss_req} (se forem lacunas reais da pessoa, registar em analise_fit.md e decidir com ela)")
    add(res, "keywords_no_topo_do_cv", any(kw_present(stop, k) for k in req) if req else True,
        "pelo menos uma keyword obrigatória no headline/resumo/competências/1º cargo")


def check_letter(letter, empresa, res, cfg):
    paras = letter.get("paragraphs", [])
    text = "\n".join(paras)
    words = len(text.split())
    lo, hi = cfg["letter_words"]
    add(res, f"carta_palavras_{lo}_{hi}", lo <= words <= hi, f"{words} palavras")
    add(res, "carta_3_a_5_paragrafos", 3 <= len(paras) <= 5, f"{len(paras)} parágrafos")
    full = json.dumps(letter, ensure_ascii=False)
    m = PLACEHOLDERS.search(full)
    add(res, "carta_sem_placeholders", not m, m.group() if m else "limpo")
    if empresa:
        add(res, "carta_menciona_empresa", norm(empresa) in norm(full), empresa)
    cl = [c for c in CLICHES if c in norm(text)]
    add(res, "carta_sem_cliches", not cl, f"clichés: {cl}", blocking=False)
    add(res, "carta_tem_assinatura_e_data", bool(letter.get("signature") and letter.get("date")), "assinatura e data")


def count_pages(path):
    """Páginas via LibreOffice+pdfinfo; None se indisponível."""
    try:
        with tempfile.TemporaryDirectory() as td:
            subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", td, str(path)],
                           capture_output=True, timeout=120, check=True)
            pdf = Path(td) / (Path(path).stem + ".pdf")
            try:
                out = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
                return int(re.search(r"Pages:\s+(\d+)", out).group(1))
            except Exception:
                data = pdf.read_bytes()
                return len(re.findall(rb"/Type\s*/Page[^s]", data)) or None
    except Exception:
        return None


def check_docx(path, cfg, words, res):
    from docx import Document
    d = Document(path)
    add(res, "docx_sem_tabelas", len(d.tables) == 0, f"{len(d.tables)} tabelas")
    add(res, "docx_sem_imagens", len(d.inline_shapes) == 0, f"{len(d.inline_shapes)} imagens")
    hf = any(p.text.strip() for s in d.sections for p in list(s.header.paragraphs) + list(s.footer.paragraphs))
    add(res, "docx_sem_texto_em_cabecalho_rodape", not hf, "cabeçalhos/rodapés ficam por vezes fora do parsing")
    pages = count_pages(path)
    if pages is None:
        est = math.ceil(words / 550)
        add(res, f"paginas<={cfg['max_pages']}", est <= cfg["max_pages"],
            f"~{est} páginas (ESTIMATIVA por nº de palavras; LibreOffice não encontrado, confirmar a abrir o ficheiro)", blocking=False)
    else:
        add(res, f"paginas<={cfg['max_pages']}", pages <= cfg["max_pages"], f"{pages} páginas")


def main():
    utf8()
    ap = argparse.ArgumentParser()
    ap.add_argument("--job", required=True, help="pasta da vaga")
    ap.add_argument("--docx", help="CV .docx a verificar (tabelas, imagens, páginas)")
    ap.add_argument("--empresa", default="")
    ap.add_argument("--round", type=int, default=None)
    ap.add_argument("--out")
    a = ap.parse_args()
    job = Path(a.job)
    cfg = load_config(job)["ats_check"]

    kwf = job / "keywords.txt"
    if not kwf.exists():
        sys.exit(f"Falta {kwf} (uma palavra-chave por linha; * = obrigatória)")
    kws = []
    for line in kwf.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line and not line.startswith("#"):
            kws.append((line.lstrip("*").strip(), line.startswith("*")))

    rnd = a.round
    if rnd is None:
        rs = [int(m.group(1)) for d in (job / "revisoes").glob("round-*") if (m := re.fullmatch(r"round-(\d+)", d.name))]
        rnd = max(rs) if rs else 1
    out = Path(a.out) if a.out else job / "revisoes" / f"round-{rnd}" / "ats_check.json"

    cv = json.loads((job / "rascunho" / "cv.json").read_text(encoding="utf-8"))
    res = {"checks": []}
    check_cv(cv, kws, res, cfg)
    letter_p = job / "rascunho" / "carta.json"
    if letter_p.exists():
        check_letter(json.loads(letter_p.read_text(encoding="utf-8")), a.empresa, res, cfg)
    if a.docx:
        check_docx(a.docx, cfg, len(cv_text(cv).split()), res)

    failed = [c for c in res["checks"] if not c["pass"] and c["blocking"]]
    res["blocking_failed"] = len(failed)
    res["warnings"] = len([c for c in res["checks"] if not c["pass"] and not c["blocking"]])
    res["status"] = "PASS" if not failed else "FAIL"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(res, ensure_ascii=False, indent=2), encoding="utf-8")
    for c in res["checks"]:
        print(("OK   " if c["pass"] else ("FAIL " if c["blocking"] else "WARN ")) + c["name"] + " - " + c["detail"])
    print("=>", res["status"], f"(gravado em {out})")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
